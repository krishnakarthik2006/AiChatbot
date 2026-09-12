"""Document loading and metadata normalisation."""
from __future__ import annotations

from pathlib import Path


SUPPORTED_SUFFIXES = {".txt", ".md", ".html", ".htm", ".pdf", ".docx", ".pptx", ".json", ".csv", ".srt", ".vtt", ".png", ".jpg", ".jpeg", ".webp"}


def read_document(path: Path) -> str:
    """Read a supported file; PDF support is intentionally optional."""
    if path.suffix.lower() == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("PDF ingestion requires pypdf. Install project requirements.") from exc
        return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    if path.suffix.lower() == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("Word ingestion requires python-docx. Install project requirements.") from exc
        return "\n".join(paragraph.text for paragraph in Document(str(path)).paragraphs)
    if path.suffix.lower() == ".pptx":
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise RuntimeError("PowerPoint ingestion requires python-pptx. Install project requirements.") from exc
        slides = []
        for index, slide in enumerate(Presentation(str(path)).slides, 1):
            text = "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip())
            if text:
                slides.append(f"Slide {index}\n{text}")
        return "\n\n".join(slides)
    if path.suffix.lower() in {".png", ".jpg", ".jpeg", ".webp"}:
        try:
            import pytesseract
            from PIL import Image
            return pytesseract.image_to_string(Image.open(path))
        except Exception as exc:
            raise RuntimeError("Image OCR requires the Tesseract desktop application to be installed and on PATH.") from exc
    return path.read_text(encoding="utf-8", errors="ignore")


def discover_documents(directory: Path) -> list[tuple[str, str]]:
    """Return non-empty local documents as (source name, body) pairs."""
    if not directory.exists():
        return []
    items: list[tuple[str, str]] = []
    for path in sorted(directory.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED_SUFFIXES:
            continue
        if path.name.lower() == "readme.md":
            continue
        body = read_document(path).strip()
        if body:
            items.append((str(path.relative_to(directory)).replace("\\", "/"), body))
    return items
